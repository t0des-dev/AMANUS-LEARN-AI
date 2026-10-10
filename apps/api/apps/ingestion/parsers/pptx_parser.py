import io
import logging
from typing import BinaryIO

import pptx

from .base import DocumentParser, InvalidDocumentFileError, ParsedDocument, ParsedPage
from .normalization import normalize_text
from .structure_detector import detect_structure

logger = logging.getLogger(__name__)

MAX_EXTRACT_SLIDES = 200
MAX_EXTRACT_CHARS = 1_000_000
MIN_SLIDE_CHARS = 15


class PPTXParser(DocumentParser):
    """Adapter for parsing Microsoft PowerPoint (.pptx) presentations.

    Preserves slide boundaries, titles, shapes, tables, speaker notes,
    and tracks empty slides and extraction warnings.
    """

    def parse(self, file_source: bytes | BinaryIO, **kwargs) -> ParsedDocument:
        if isinstance(file_source, bytes):
            stream = io.BytesIO(file_source)
        else:
            stream = file_source

        try:
            prs = pptx.Presentation(stream)
        except Exception as e:
            raise InvalidDocumentFileError(f"Fichier PPTX invalide ou corrompu : {e}") from e

        global_metadata: dict = {"warnings": []}
        warnings: list[str] = global_metadata["warnings"]

        try:
            core_props = prs.core_properties
            if core_props:
                if core_props.title:
                    global_metadata["title"] = core_props.title
                if core_props.author:
                    global_metadata["author"] = core_props.author
        except Exception as e:
            logger.debug(f"[PPTXParser] Error extracting core properties: {e}")

        total_slides = len(prs.slides)
        if total_slides == 0:
            raise InvalidDocumentFileError("La présentation PPTX ne contient aucune diapositive.")

        slides_to_process = total_slides
        if total_slides > MAX_EXTRACT_SLIDES:
            warnings.append(f"MAX_SLIDES_LIMIT_REACHED_{MAX_EXTRACT_SLIDES}_OF_{total_slides}")
            slides_to_process = MAX_EXTRACT_SLIDES

        pages: list[ParsedPage] = []
        accumulated_chars = 0

        for idx in range(slides_to_process):
            slide = prs.slides[idx]
            page_num = idx + 1
            slide_elements: list[str] = []
            slide_title: str | None = None
            page_warnings: list[str] = []

            # Check title placeholder if available
            try:
                if slide.shapes.title and slide.shapes.title.text.strip():
                    slide_title = slide.shapes.title.text.strip()
                    slide_elements.append(f"# {slide_title}")
            except Exception:
                pass

            # Extract text from shapes and tables
            for shape in slide.shapes:
                try:
                    # Avoid duplicating slide title
                    if slide_title and hasattr(shape, "text") and shape.text.strip() == slide_title:
                        continue

                    if shape.has_text_frame:
                        for paragraph in shape.text_frame.paragraphs:
                            text = paragraph.text.strip()
                            if text:
                                slide_elements.append(text)

                    elif shape.has_table:
                        table_rows: list[str] = []
                        for row in shape.table.rows:
                            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                            if cells:
                                table_rows.append(" | ".join(cells))
                        if table_rows:
                            slide_elements.append("[Tableau]\n" + "\n".join(table_rows))
                except Exception as e:
                    logger.warning(f"[PPTXParser] Error reading shape on slide {page_num}: {e}")
                    page_warnings.append(f"SLIDE_{page_num}_SHAPE_READ_ERROR")

            # Extract speaker notes if any
            try:
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                    notes = slide.notes_slide.notes_text_frame.text.strip()
                    if notes:
                        slide_elements.append(f"[Notes de présentation]\n{notes}")
            except Exception:
                pass

            raw_slide_text = "\n\n".join(slide_elements)
            clean_text = normalize_text(raw_slide_text)

            if len(clean_text) < MIN_SLIDE_CHARS:
                page_warnings.append(f"SLIDE_{page_num}_EMPTY")

            accumulated_chars += len(clean_text)
            struct_info = detect_structure(clean_text)
            page_meta = {
                "page_number": page_num,
                "slide_title": slide_title,
                "chapter": struct_info.get("chapter")
                or (f"Diapositive {page_num}" if not slide_title else None),
                "section": struct_info.get("section") or slide_title,
                "subsection": struct_info.get("subsection"),
                "char_count": len(clean_text),
                "headings": struct_info.get("headings", []),
                "is_empty": len(clean_text) < MIN_SLIDE_CHARS,
                "warnings": page_warnings,
            }

            pages.append(
                ParsedPage(
                    page_number=page_num,
                    text=clean_text,
                    ocr_used=False,
                    metadata=page_meta,
                )
            )

            if accumulated_chars >= MAX_EXTRACT_CHARS:
                warnings.append("MAX_CHARS_LIMIT_REACHED")
                break

        return ParsedDocument(pages=pages, metadata=global_metadata)
