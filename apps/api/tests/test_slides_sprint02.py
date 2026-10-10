import io

import pytest
from pptx import Presentation as PptxPresentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches
from rest_framework import status
from rest_framework.test import APIClient

from apps.accounts.models import User
from apps.courses.models import Course, CourseSection
from apps.organizations.models import Organization, OrganizationMember, RoleChoices
from apps.slides.models import (
    Presentation,
    PresentationSlide,
    PresentationTheme,
)
from apps.slides.services.pptx_exporter import (
    PPTXExporter,
    detect_slide_layout,
    get_font_scale_for_bullets,
    is_arabic_text,
)
from apps.slides.services.slide_generator import SlideGenerator
from apps.slides.services.slide_planner import SlidePlanner
from apps.slides.services.validator import (
    InvalidPresentationPayloadError,
    PresentationValidator,
)


@pytest.fixture
def org_and_user(db):
    org = Organization.objects.create(name="Amanus Academy", slug="amanus-academy-s2")
    user = User.objects.create_user(
        email="instructor-s2@amanus.test",
        password="TestPassword123!",
        first_name="Professeur",
        last_name="Sprint02",
    )
    OrganizationMember.objects.create(organization=org, user=user, role=RoleChoices.ADMIN)
    return org, user


@pytest.fixture
def course_with_sections(db, org_and_user):
    org, user = org_and_user
    course = Course.objects.create(
        organization=org,
        created_by=user,
        title="Apprentissage Profond & Architectures Modernes",
        language="fr",
        level="INTERMEDIATE",
    )

    sec1 = CourseSection.objects.create(
        course=course,
        title="Fondements des Réseaux Neuronaux",
        summary="Comprendre le fonctionnement des neurones artificiels et rétropropagation.",
        order=1,
    )
    sec2 = CourseSection.objects.create(
        course=course,
        title="Architectures Transformers et Attention",
        summary="Mécanismes d'attention multi-têtes et traitement séquentiel parallèle.",
        order=2,
    )
    return course, [sec1, sec2]


@pytest.mark.django_db
class TestSprint02SlidePlannerAndPedagogy:
    """Tests multilingual pedagogical planning, level adaptation, and layout generation."""

    def test_multilingual_slide_planner_fr(self, course_with_sections):
        course, _ = course_with_sections
        planner = SlidePlanner()
        plan = planner.plan_presentation(course, language="fr", level="ADVANCED")

        assert len(plan) >= 4
        # Title slide
        assert plan[0]["slide_type"] == "title"
        assert "Formation" in plan[0]["content"]
        assert "ADVANCED" in plan[0]["content"]

        # Agenda slide
        assert plan[1]["slide_type"] == "agenda"
        assert "Sommaire" in plan[1]["title"]
        assert "Module 1" in plan[1]["content"]

        # Summary slide
        summary_slide = [s for s in plan if s["slide_type"] == "summary"][0]
        assert "Synthèse" in summary_slide["title"]

        # Conclusion slide
        conclusion_slide = [s for s in plan if s["slide_type"] == "conclusion"][0]
        assert "Conclusion" in conclusion_slide["title"]

    def test_multilingual_slide_planner_ar(self, course_with_sections):
        course, _ = course_with_sections
        planner = SlidePlanner()
        plan = planner.plan_presentation(
            course, language="ar", level="BEGINNER", focus="التطبيقات الميدانية"
        )

        assert len(plan) >= 4
        # Title slide in Arabic
        assert plan[0]["slide_type"] == "title"
        assert "المقرر التعليمي" in plan[0]["content"]

        # Agenda slide in Arabic
        assert plan[1]["slide_type"] == "agenda"
        assert "جدول الأعمال" in plan[1]["title"]
        assert "المحور 1" in plan[1]["content"]

        # Summary slide with focus
        summary_slide = [s for s in plan if s["slide_type"] == "summary"][0]
        assert "الخلاصة" in summary_slide["title"]
        assert "التطبيقات الميدانية" in summary_slide["content"]

        # Conclusion slide
        conclusion_slide = [s for s in plan if s["slide_type"] == "conclusion"][0]
        assert "الخاتمة" in conclusion_slide["title"]

    def test_multilingual_slide_planner_en(self, course_with_sections):
        course, _ = course_with_sections
        planner = SlidePlanner()
        plan = planner.plan_presentation(course, language="en", level="EXPERT")

        assert len(plan) >= 4
        assert plan[0]["slide_type"] == "title"
        assert "Course" in plan[0]["content"]
        assert plan[1]["slide_type"] == "agenda"
        assert "Agenda" in plan[1]["title"]
        assert "Module 1" in plan[1]["content"]

    def test_speaker_notes_and_image_prompts_generated(self, course_with_sections):
        course, _ = course_with_sections
        planner = SlidePlanner()
        plan = planner.plan_presentation(course, language="fr")

        for slide_data in plan:
            assert "speaker_notes" in slide_data
            assert len(slide_data["speaker_notes"]) > 10
            assert "image_prompt" in slide_data
            assert len(slide_data["image_prompt"]) > 10


@pytest.mark.django_db
class TestSprint02PresentationValidator:
    """Tests text sanitization, anti-overflow limits, and deck validation."""

    def test_sanitize_text_strips_control_characters(self):
        corrupted = "Titre\x00avec\x08caractères\x1bnon-imprimables"
        cleaned = PresentationValidator.sanitize_text(corrupted)
        assert cleaned == "Titreaveccaractèresnon-imprimables"

    def test_validate_presentation_metadata_rejects_empty_title(self):
        with pytest.raises(InvalidPresentationPayloadError):
            PresentationValidator.validate_presentation_metadata("   ", "modern_dark")

    def test_validate_slide_content_truncates_excessive_bullets(self):
        content = "\n".join([f"• Point numéro {i}" for i in range(1, 12)])
        validated = PresentationValidator.validate_slide_content("Mon Titre", content, 3)

        lines = [line for line in validated["content"].split("\n") if line.strip()]
        assert len(lines) == PresentationValidator.MAX_BULLETS_PER_SLIDE
        assert len(lines) <= 6

    def test_validate_slide_content_truncates_long_bullet(self):
        long_line = "A" * 200
        validated = PresentationValidator.validate_slide_content("Titre", long_line, 1)
        assert len(validated["content"]) <= PresentationValidator.MAX_CHARS_PER_BULLET + 10
        assert validated["content"].endswith("...")

    def test_validate_deck_for_export_rejects_zero_slides(self, course_with_sections):
        course, _ = course_with_sections
        empty_pres = Presentation.objects.create(
            course=course,
            title="Présentation Vide",
            theme=PresentationTheme.MODERN_DARK,
        )

        with pytest.raises(InvalidPresentationPayloadError):
            PresentationValidator.validate_deck_for_export(empty_pres)


@pytest.mark.django_db
class TestSprint02PPTXExporter:
    """Tests PPTX generation, 16:9 dimensions, reopening, RTL, and layouts."""

    def test_pptx_exporter_16_9_widescreen_and_reopening(self, course_with_sections):
        course, _ = course_with_sections
        generator = SlideGenerator()
        pres = generator.generate_presentation(course, title="Test Deck 16:9")

        exporter = PPTXExporter()
        buffer = exporter.export_to_buffer(pres)
        assert buffer.getvalue()[:4] == b"PK\x03\x04"  # Valid ZIP / OpenXML header

        # Structural Reopening Test with python-pptx
        reopened = PptxPresentation(io.BytesIO(buffer.getvalue()))
        assert len(reopened.slides) == pres.slides.count()

        # Check 16:9 widescreen dimensions
        assert round(reopened.slide_width / Inches(1), 3) == 13.333
        assert round(reopened.slide_height / Inches(1), 3) == 7.5

    def test_pptx_detect_slide_layout(self):
        assert detect_slide_layout(1, 5, "Titre Général", "") == "title"
        assert detect_slide_layout(2, 5, "Sommaire & Objectifs", "") == "agenda"
        assert detect_slide_layout(3, 5, "Architecture", "Colonne 1 || Colonne 2") == "two_column"
        assert detect_slide_layout(4, 5, "Synthèse & Bilan", "Résumé") == "summary"
        assert detect_slide_layout(5, 5, "Conclusion & Questions", "Merci") == "conclusion"
        assert detect_slide_layout(3, 5, "Notions Clés", "Contenu classique") == "concept"

    def test_pptx_two_column_slide_export_and_reopen(self, course_with_sections):
        course, _ = course_with_sections
        pres = Presentation.objects.create(
            course=course,
            title="Comparatif Algorithmes",
            theme=PresentationTheme.CORPORATE_BLUE,
        )
        PresentationSlide.objects.create(
            presentation=pres,
            slide_number=1,
            title="Comparatif Algorithmes",
            content="• Module 1\n• Module 2",
        )
        PresentationSlide.objects.create(
            presentation=pres,
            slide_number=2,
            title="Analyse Comparative Méthodes",
            content="• Approche Supervisée : Données étiquetées\n• Précision élevée\n || \n• Approche Non-Supervisée : Données brutes\n• Découverte de clusters",
            speaker_notes="Comparer les deux approches",
        )

        exporter = PPTXExporter()
        buffer = exporter.export_to_buffer(pres)
        reopened = PptxPresentation(io.BytesIO(buffer.getvalue()))

        assert len(reopened.slides) == 2
        slide2 = reopened.slides[1]
        # Slide 2 should have left and right card shapes plus text boxes
        assert len(slide2.shapes) >= 4
        # Notes check
        assert slide2.notes_slide.notes_text_frame.text == "Comparer les deux approches"

    def test_pptx_arabic_rtl_alignment_and_unicode_reopen(self, course_with_sections):
        course, _ = course_with_sections
        pres = Presentation.objects.create(
            course=course,
            title="الذكاء الاصطناعي التوليدي",
            theme=PresentationTheme.ACADEMIC_INDIGO,
        )
        PresentationSlide.objects.create(
            presentation=pres,
            slide_number=1,
            title="الذكاء الاصطناعي التوليدي",
            content="• المقرر : تعلم الآلة المتقدم\n• المستوى : متوسط",
        )
        PresentationSlide.objects.create(
            presentation=pres,
            slide_number=2,
            title="نماذج اللغات الضخمة وتطبيقاتها",
            content="• فهم آليات الانتباه الذاتي\n• التطبيق العملي في معالجة اللغات الطبيعية\n• التقييم المعياري للنماذج",
            speaker_notes="شرح مفصل حول المعالجة الآلية",
        )

        assert is_arabic_text("نماذج اللغات الضخمة") is True

        exporter = PPTXExporter()
        buffer = exporter.export_to_buffer(pres)
        reopened = PptxPresentation(io.BytesIO(buffer.getvalue()))

        assert len(reopened.slides) == 2
        slide2 = reopened.slides[1]

        # Check footer text contains Arabic slide label
        found_arabic_footer = False
        for shape in slide2.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    if "الشريحة 2 من 2" in p.text:
                        found_arabic_footer = True
                        assert p.alignment == PP_ALIGN.RIGHT
        assert found_arabic_footer is True

    def test_pptx_font_scaling_anti_overflow(self):
        sz_small, sp_small = get_font_scale_for_bullets(2)
        sz_med, sp_med = get_font_scale_for_bullets(4)
        sz_large, sp_large = get_font_scale_for_bullets(6)

        assert sz_small > sz_med > sz_large
        assert sp_small > sp_med > sp_large

    def test_pptx_preserves_teacher_manual_edits(self, course_with_sections):
        course, _ = course_with_sections
        generator = SlideGenerator()
        pres = generator.generate_presentation(course, title="Deck Original")

        # Teacher modifies slide 2
        slide2 = pres.slides.get(slide_number=2)
        slide2.title = "Titre Modifié Manuellement par l'Enseignant"
        slide2.content = "• Contenu personnalisé 1\n• Contenu personnalisé 2"
        slide2.save()

        exporter = PPTXExporter()
        buffer = exporter.export_to_buffer(pres)
        reopened = PptxPresentation(io.BytesIO(buffer.getvalue()))

        # Verify teacher edits are present in the re-exported PPTX
        slide2_reopened = reopened.slides[1]
        titles_in_slide = [
            shape.text_frame.text
            for shape in slide2_reopened.shapes
            if shape.has_text_frame and "Titre Modifié Manuellement" in shape.text_frame.text
        ]
        assert len(titles_in_slide) > 0


@pytest.mark.django_db
class TestSprint02APIEndpoints:
    """Tests API presentation creation with options and export validation."""

    def test_api_generate_presentation_with_language_level_focus(
        self, org_and_user, course_with_sections
    ):
        _, user = org_and_user
        course, _ = course_with_sections
        client = APIClient()
        client.force_authenticate(user=user)

        response = client.post(
            f"/api/v1/courses/{course.id}/presentations/",
            {
                "title": "عرض تقديمي باللغة العربية",
                "theme": "academic_indigo",
                "language": "ar",
                "level": "ADVANCED",
                "focus": "التعلم العميق التوليدي",
            },
            format="json",
        )

        assert response.status_code == status.HTTP_201_CREATED
        data = response.data
        assert data["title"] == "عرض تقديمي باللغة العربية"
        assert data["theme"] == "academic_indigo"
        assert len(data["slides"]) >= 4

        # Verify Arabic agenda title in slides
        slide_titles = [s["title"] for s in data["slides"]]
        assert any("جدول الأعمال" in t for t in slide_titles)

    def test_api_export_rejects_empty_presentation(self, org_and_user, course_with_sections):
        _, user = org_and_user
        course, _ = course_with_sections
        client = APIClient()
        client.force_authenticate(user=user)

        empty_pres = Presentation.objects.create(
            course=course,
            title="Présentation Sans Diapositives",
            theme=PresentationTheme.MODERN_DARK,
        )

        response = client.post(f"/api/v1/presentations/{empty_pres.id}/export/")
        assert response.status_code == status.HTTP_400_BAD_REQUEST
        assert "empty_presentation" in response.data.get("code", "")
