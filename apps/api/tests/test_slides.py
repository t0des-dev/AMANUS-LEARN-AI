import io
import shutil
import tempfile
import uuid
from unittest.mock import MagicMock, patch

from django.contrib.auth import get_user_model
from pptx import Presentation as PptxPresentation
from rest_framework import status
from rest_framework.test import APITestCase
from rest_framework_simplejwt.tokens import RefreshToken

from apps.courses.models import Course, CourseLevel, CourseSection, CourseStatus
from apps.documents.services.storage import LocalStorageService, reset_storage_service
from apps.organizations.models import Organization, OrganizationMember, RoleChoices
from apps.slides.models import (
    Presentation,
    PresentationSlide,
    PresentationStatus,
    PresentationTheme,
)
from apps.slides.services.pptx_exporter import PPTXExporter
from apps.slides.services.slide_generator import SlideGenerator
from apps.slides.services.slide_planner import SlidePlanner
from apps.slides.tasks import export_presentation_task

User = get_user_model()


class SlidePlannerAndGeneratorTests(APITestCase):
    """Unit tests for SlidePlanner and SlideGenerator services."""

    def setUp(self):
        self.user = User.objects.create_user(
            email="teacher@amanus.test",
            password="testpassword123",
            first_name="Prof",
            last_name="Martin",
        )
        self.org = Organization.objects.create(name="Univ Tech", slug="univ-tech")
        OrganizationMember.objects.create(
            organization=self.org,
            user=self.user,
            role=RoleChoices.TEACHER,
        )
        self.course = Course.objects.create(
            organization=self.org,
            created_by=self.user,
            title="Introduction au Deep Learning",
            level=CourseLevel.INTERMEDIATE,
            status=CourseStatus.PUBLISHED,
        )
        # Create sections
        self.sec1 = CourseSection.objects.create(
            course=self.course,
            title="Les Perceptrons et Neurones Artificiels",
            order=1,
            summary="Historique et fondements mathématiques des fonctions d'activation.",
            content="• Fonction sigmoïde et ReLU\n• Modèle de Rosenblatt\n• Propagation avant et calcul matriciel",
        )
        self.sec2 = CourseSection.objects.create(
            course=self.course,
            title="Rétropropagation et Descente de Gradient",
            order=2,
            summary="Algorithme d'optimisation et calcul des gradients partiels.",
            content="• Fonction de coût MSE et Cross-Entropy\n• Taux d'apprentissage et momentum\n• Algorithme Adam",
        )

    def test_slide_planner_generates_pedagogical_blueprints(self):
        planner = SlidePlanner()
        blueprints = planner.plan_presentation(self.course)

        self.assertGreaterEqual(len(blueprints), 4)
        # First slide is Title
        self.assertEqual(blueprints[0]["slide_type"], "title")
        self.assertIn("Deep Learning", blueprints[0]["title"])
        self.assertIn("Univ Tech", blueprints[0]["content"])

        # Second slide is Agenda
        self.assertEqual(blueprints[1]["slide_type"], "agenda")
        self.assertIn("Sommaire", blueprints[1]["title"])

        # Concept slides for sections
        concept_titles = [bp["title"] for bp in blueprints if bp["slide_type"] == "concept"]
        self.assertIn(self.sec1.title, concept_titles)
        self.assertIn(self.sec2.title, concept_titles)

        # Last slides are summary and conclusion
        self.assertEqual(blueprints[-2]["slide_type"], "summary")
        self.assertEqual(blueprints[-1]["slide_type"], "conclusion")

        # Check speaker notes and image prompt presence
        for bp in blueprints:
            self.assertTrue(len(bp.get("speaker_notes", "")) > 10)
            self.assertTrue(len(bp.get("image_prompt", "")) > 5)

    def test_slide_generator_creates_presentation_and_slides(self):
        generator = SlideGenerator()
        presentation = generator.generate_presentation(
            course=self.course,
            title="Présentation Deep Learning Interactive",
            theme=PresentationTheme.ACADEMIC_INDIGO,
        )

        self.assertIsInstance(presentation, Presentation)
        self.assertEqual(presentation.status, PresentationStatus.READY)
        self.assertEqual(presentation.title, "Présentation Deep Learning Interactive")
        self.assertEqual(presentation.theme, PresentationTheme.ACADEMIC_INDIGO)

        slides = presentation.slides.all().order_by("slide_number")
        self.assertGreaterEqual(slides.count(), 4)

        # Check order numbers
        orders = [s.slide_number for s in slides]
        self.assertEqual(orders, list(range(1, len(orders) + 1)))

        # Check slide content
        first_slide = slides.first()
        self.assertEqual(first_slide.title, "Présentation Deep Learning Interactive")
        self.assertTrue(first_slide.speaker_notes)


class PPTXExporterTests(APITestCase):
    """Unit tests for PPTX export formatting and storage."""

    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.user = User.objects.create_user(
            email="export_teacher@amanus.test",
            password="testpassword123",
        )
        self.org = Organization.objects.create(name="Slide Org", slug="slide-org")
        OrganizationMember.objects.create(
            organization=self.org, user=self.user, role=RoleChoices.TEACHER
        )
        self.course = Course.objects.create(
            organization=self.org,
            created_by=self.user,
            title="Architecture des Systèmes Cloud",
        )
        generator = SlideGenerator()
        self.presentation = generator.generate_presentation(
            course=self.course,
            title="Architecture Cloud Deck",
            theme=PresentationTheme.MODERN_DARK,
        )

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)
        reset_storage_service()

    def test_pptx_exporter_generates_valid_16_9_presentation(self):
        exporter = PPTXExporter()
        buf = exporter.export_to_buffer(self.presentation)

        self.assertIsInstance(buf, io.BytesIO)
        self.assertGreater(buf.getbuffer().nbytes, 1000)

        # Parse generated pptx with python-pptx to verify integrity
        prs = PptxPresentation(buf)
        # Check widescreen 16:9 dimensions (13.333 x 7.5 inches)
        self.assertAlmostEqual(prs.slide_width.inches, 13.333, places=2)
        self.assertAlmostEqual(prs.slide_height.inches, 7.5, places=2)

        # Check slides count matches database
        self.assertEqual(len(prs.slides), self.presentation.slides.count())

        # Check speaker notes on slide
        slide_1 = prs.slides[0]
        notes_text = slide_1.notes_slide.notes_text_frame.text
        self.assertTrue(len(notes_text) > 0)

    def test_pptx_export_and_save_with_storage(self):
        storage_service = LocalStorageService(base_dir=self.temp_dir)
        exporter = PPTXExporter(storage_service=storage_service)

        storage_key = exporter.export_and_save(self.presentation)
        self.assertTrue(storage_key.endswith(".pptx"))
        self.assertEqual(self.presentation.status, PresentationStatus.READY)
        self.assertEqual(self.presentation.storage_key, storage_key)
        self.assertTrue(storage_service.file_exists(storage_key))

    def test_async_celery_export_task(self):
        with patch("apps.slides.tasks.PPTXExporter") as mock_exporter_cls:
            mock_inst = MagicMock()
            mock_inst.export_and_save.return_value = "presentations/test/pres.pptx"
            mock_exporter_cls.return_value = mock_inst

            res = export_presentation_task(str(self.presentation.id))
            self.assertEqual(res["status"], "READY")
            self.assertEqual(res["presentation_id"], str(self.presentation.id))


class PresentationAPITests(APITestCase):
    """Integration tests for all Presentation and Slide API endpoints."""

    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.storage_service = LocalStorageService(base_dir=self.temp_dir)

        self.user = User.objects.create_user(
            email="author@amanus.test",
            password="testpassword123",
            first_name="Alice",
            last_name="Prof",
        )
        self.student = User.objects.create_user(
            email="student@amanus.test",
            password="testpassword123",
            first_name="Bob",
            last_name="Learner",
        )
        self.stranger = User.objects.create_user(
            email="stranger@amanus.test",
            password="testpassword123",
        )

        self.org = Organization.objects.create(name="AI Academy", slug="ai-academy")
        OrganizationMember.objects.create(
            organization=self.org,
            user=self.user,
            role=RoleChoices.TEACHER,
        )
        OrganizationMember.objects.create(
            organization=self.org,
            user=self.student,
            role=RoleChoices.STUDENT,
        )

        self.course = Course.objects.create(
            organization=self.org,
            created_by=self.user,
            title="Algorithmes Fondamentaux",
            level=CourseLevel.BEGINNER,
        )
        CourseSection.objects.create(
            course=self.course,
            title="Complexité et Notations Big-O",
            order=1,
            summary="Analyse de la complexité spatiale et temporelle des algorithmes.",
        )

        token = RefreshToken.for_user(self.user)
        self.auth_headers = {"HTTP_AUTHORIZATION": f"Bearer {token.access_token}"}

        student_token = RefreshToken.for_user(self.student)
        self.student_headers = {"HTTP_AUTHORIZATION": f"Bearer {student_token.access_token}"}

        stranger_token = RefreshToken.for_user(self.stranger)
        self.stranger_headers = {"HTTP_AUTHORIZATION": f"Bearer {stranger_token.access_token}"}

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)
        reset_storage_service()

    def test_generate_presentation_from_course(self):
        url = f"/api/v1/courses/{self.course.id}/presentations/"
        payload = {
            "title": "Algorithmes & Structures",
            "theme": PresentationTheme.CORPORATE_BLUE,
        }
        res = self.client.post(url, payload, format="json", **self.auth_headers)

        self.assertEqual(res.status_code, status.HTTP_201_CREATED)
        self.assertEqual(res.data["title"], "Algorithmes & Structures")
        self.assertEqual(res.data["theme"], PresentationTheme.CORPORATE_BLUE)
        self.assertEqual(res.data["status"], PresentationStatus.READY)
        self.assertGreaterEqual(len(res.data["slides"]), 4)

    def test_list_presentations_for_course(self):
        # Generate one presentation
        gen = SlideGenerator()
        gen.generate_presentation(course=self.course)

        url = f"/api/v1/courses/{self.course.id}/presentations/"
        res = self.client.get(url, **self.auth_headers)

        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(len(res.data), 1)
        self.assertEqual(res.data[0]["course_title"], self.course.title)

    def test_retrieve_presentation_detail(self):
        gen = SlideGenerator()
        pres = gen.generate_presentation(course=self.course)

        url = f"/api/v1/presentations/{pres.id}/"
        res = self.client.get(url, **self.auth_headers)

        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(res.data["id"], str(pres.id))
        self.assertIsInstance(res.data["slides"], list)
        self.assertGreaterEqual(len(res.data["slides"]), 4)

    def test_patch_presentation_title_and_theme(self):
        gen = SlideGenerator()
        pres = gen.generate_presentation(course=self.course)

        url = f"/api/v1/presentations/{pres.id}/"
        payload = {
            "title": "Nouveau Titre Modifié",
            "theme": PresentationTheme.MINIMAL_LIGHT,
        }
        res = self.client.patch(url, payload, format="json", **self.auth_headers)

        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(res.data["title"], "Nouveau Titre Modifié")
        self.assertEqual(res.data["theme"], PresentationTheme.MINIMAL_LIGHT)

    def test_add_slide_to_presentation(self):
        gen = SlideGenerator()
        pres = gen.generate_presentation(course=self.course)
        initial_count = pres.slides.count()

        url = f"/api/v1/presentations/{pres.id}/slides/"
        payload = {
            "title": "Exercice Pratique 1",
            "content": "• Implémenter le tri fusion en Python\n• Mesurer le temps d'exécution",
            "speaker_notes": "Laissez 10 minutes aux apprenants pour coder l'exercice.",
        }
        res = self.client.post(url, payload, format="json", **self.auth_headers)

        self.assertEqual(res.status_code, status.HTTP_201_CREATED)
        self.assertEqual(res.data["title"], "Exercice Pratique 1")
        self.assertEqual(res.data["slide_number"], initial_count + 1)
        self.assertEqual(pres.slides.count(), initial_count + 1)

    def test_edit_slide_content(self):
        gen = SlideGenerator()
        pres = gen.generate_presentation(course=self.course)
        slide = pres.slides.first()

        url = f"/api/v1/presentations/{pres.id}/slides/{slide.id}/"
        payload = {
            "title": "Titre Slide Modifié",
            "content": "• Point clé 1\n• Point clé 2",
            "speaker_notes": "Nouvelles notes orateur",
        }
        res = self.client.patch(url, payload, format="json", **self.auth_headers)

        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(res.data["title"], "Titre Slide Modifié")
        slide.refresh_from_db()
        self.assertEqual(slide.title, "Titre Slide Modifié")
        self.assertEqual(slide.speaker_notes, "Nouvelles notes orateur")

    def test_delete_slide_and_verify_renumbering(self):
        gen = SlideGenerator()
        pres = gen.generate_presentation(course=self.course)
        slides = list(pres.slides.all().order_by("slide_number"))
        slide_to_delete = slides[1]  # Slide 2

        url = f"/api/v1/presentations/{pres.id}/slides/{slide_to_delete.id}/"
        res = self.client.delete(url, **self.auth_headers)

        self.assertEqual(res.status_code, status.HTTP_204_NO_CONTENT)
        self.assertFalse(PresentationSlide.objects.filter(id=slide_to_delete.id).exists())

        # Verify no gap in slide_numbers
        remaining_numbers = list(
            pres.slides.all().order_by("slide_number").values_list("slide_number", flat=True)
        )
        self.assertEqual(remaining_numbers, list(range(1, len(remaining_numbers) + 1)))

    def test_reorder_slides(self):
        gen = SlideGenerator()
        pres = gen.generate_presentation(course=self.course)
        slides = list(pres.slides.all().order_by("slide_number"))

        url = f"/api/v1/presentations/{pres.id}/reorder/"
        # Invert slide 1 and slide 2
        payload = {
            "slides": [
                {"id": str(slides[0].id), "slide_number": 2},
                {"id": str(slides[1].id), "slide_number": 1},
            ]
        }
        res = self.client.post(url, payload, format="json", **self.auth_headers)

        self.assertEqual(res.status_code, status.HTTP_200_OK)
        slides[0].refresh_from_db()
        slides[1].refresh_from_db()
        self.assertEqual(slides[0].slide_number, 2)
        self.assertEqual(slides[1].slide_number, 1)

    @patch("apps.slides.services.pptx_exporter.get_storage_service")
    def test_export_presentation_pptx_sync(self, mock_storage_factory):
        mock_storage_factory.return_value = self.storage_service
        gen = SlideGenerator()
        pres = gen.generate_presentation(course=self.course)

        url = f"/api/v1/presentations/{pres.id}/export/"
        res = self.client.post(url, {}, format="json", **self.auth_headers)

        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(res.data["status"], PresentationStatus.READY)
        self.assertTrue(res.data["storage_key"].endswith(".pptx"))
        self.assertIsNotNone(res.data["export_url"])

    @patch("apps.slides.views.export_presentation_task.delay")
    def test_export_presentation_pptx_async(self, mock_celery_delay):
        mock_task = MagicMock()
        mock_task.id = uuid.uuid4()
        mock_celery_delay.return_value = mock_task

        gen = SlideGenerator()
        pres = gen.generate_presentation(course=self.course)

        url = f"/api/v1/presentations/{pres.id}/export/?async=true"
        res = self.client.post(url, {}, format="json", **self.auth_headers)

        self.assertEqual(res.status_code, status.HTTP_202_ACCEPTED)
        self.assertEqual(res.data["status"], PresentationStatus.EXPORTING)
        self.assertEqual(res.data["task_id"], str(mock_task.id))
        mock_celery_delay.assert_called_once_with(str(pres.id))

    def test_permission_denied_for_non_member(self):
        gen = SlideGenerator()
        pres = gen.generate_presentation(course=self.course)

        url = f"/api/v1/presentations/{pres.id}/"
        res = self.client.get(url, **self.stranger_headers)
        self.assertEqual(res.status_code, status.HTTP_403_FORBIDDEN)

    def test_student_read_only_permission(self):
        gen = SlideGenerator()
        pres = gen.generate_presentation(course=self.course)

        # Student can read
        url = f"/api/v1/presentations/{pres.id}/"
        res = self.client.get(url, **self.student_headers)
        self.assertEqual(res.status_code, status.HTTP_200_OK)

        # Student cannot edit
        res_patch = self.client.patch(
            url, {"title": "Hacked"}, format="json", **self.student_headers
        )
        self.assertEqual(res_patch.status_code, status.HTTP_403_FORBIDDEN)
